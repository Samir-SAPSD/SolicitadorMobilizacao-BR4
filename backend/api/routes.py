import json
import os
import re
import subprocess
import threading
from datetime import datetime

import openpyxl
from flask import Blueprint, Response, current_app, jsonify, render_template, request, send_from_directory, stream_with_context
from werkzeug.utils import secure_filename

from backend.config import SCRIPTS_DIR
from backend.state import end_job, start_job
from backend.services import (
    TEMPLATE_FILENAME,
    allowed_file,
    build_powershell_command,
    build_powershell_populate_command,
    build_powershell_template_update_command,
    decode_powershell_output,
    extract_lista_suspensa_columns,
    generate_report_excel,
    get_template_status,
    normalize_key,
    _create_group_workbook,
    _read_grouped_excel,
    _popen_hidden_kwargs,
    repair_mojibake,
    save_template_status,
)

api_bp = Blueprint('api', __name__)


@api_bp.route('/')
def index():
    return render_template('index.html')


@api_bp.route('/download-template')
def download_template():
    return send_from_directory(current_app.config['UPLOAD_FOLDER'], TEMPLATE_FILENAME, as_attachment=True)


@api_bp.route('/template-update-status', methods=['GET'])
def template_update_status():
    return jsonify(get_template_status()), 200


@api_bp.route('/update-template', methods=['POST'])
def update_template():
    start_job()
    try:
        template_path = os.path.abspath(os.path.join(current_app.config['UPLOAD_FOLDER'], TEMPLATE_FILENAME))
        if not os.path.exists(template_path):
            return jsonify({'status': 'error', 'message': 'Template não encontrado.'}), 404

        script_path = os.path.join(str(SCRIPTS_DIR), 'Update-ExcelTemplateChoices.ps1')
        if not os.path.exists(script_path):
            return jsonify({'status': 'error', 'message': 'Script de atualização não encontrado.'}), 500

        cmd = build_powershell_template_update_command(script_path, template_path)

        process = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            universal_newlines=False,
            **_popen_hidden_kwargs()
        )
        raw_output = process.communicate()[0]
        process.wait()
        full_output = decode_powershell_output(raw_output)

        if process.returncode != 0:
            return jsonify({
                'status': 'error',
                'message': 'Falha ao atualizar template.',
                'log': full_output
            }), 500

        updated_at = datetime.now().isoformat()
        save_template_status(updated_at)

        return jsonify({
            'status': 'success',
            'message': 'Template atualizado com sucesso.',
            'last_updated': updated_at,
            'log': full_output
        }), 200
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'Erro ao atualizar template: {str(e)}'
        }), 500
    finally:
        end_job()


@api_bp.route('/validate', methods=['POST'])
def validate():
    """Fase 1: valida estrutura e dados por GRUPO antes do upload."""
    start_job()
    try:
        file = request.files.get('file')

        if not file or not allowed_file(file.filename):
            return jsonify({'status': 'error', 'errors': ['Arquivo inválido ou não enviado.']}), 400

        try:
            file.stream.seek(0, os.SEEK_END)
            uploaded_size = file.stream.tell()
            file.stream.seek(0)
        except Exception:
            uploaded_size = None

        if uploaded_size == 0:
            return jsonify({'status': 'error', 'errors': ['O arquivo enviado está vazio.']}), 400

        filename = secure_filename(file.filename)
        if not filename:
            return jsonify({'status': 'error', 'errors': ['Nome de arquivo inválido.']}), 400

        file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)

        if os.path.getsize(file_path) == 0:
            if os.path.exists(file_path):
                os.remove(file_path)
            return jsonify({'status': 'error', 'errors': ['O arquivo enviado está vazio.']}), 400

        grouped_data, grouping_errors = _read_grouped_excel(file_path)
        if grouping_errors:
            return jsonify({
                'status': 'failed',
                'errors': grouping_errors,
                'filename': filename,
                'total_lines': 0
            }), 200

        script_path = os.path.join(str(SCRIPTS_DIR), 'Validate-ExcelData.ps1')
        if not os.path.exists(script_path):
            return jsonify({
                'status': 'error',
                'errors': ['Script de validação não encontrado.'],
                'filename': filename
            }), 500

        all_errors: list[str] = []
        all_logs: list[str] = []
        total_people = 0
        total_equip = 0

        group_summary = []
        for group_key in grouped_data['group_order']:
            people_count = len(grouped_data['sheets']['PESSOAS']['rows_by_group'].get(group_key, []))
            equip_count = len(grouped_data['sheets']['EQUIPAMENTOS']['rows_by_group'].get(group_key, []))
            total_people += people_count
            total_equip += equip_count
            group_summary.append({
                'group': group_key,
                'qtd_pessoas': people_count,
                'qtd_equipamentos': equip_count
            })

            temp_group_path = _create_group_workbook(grouped_data, group_key)
            try:
                for sheet_name in ('PESSOAS', 'EQUIPAMENTOS'):
                    row_count = len(grouped_data['sheets'][sheet_name]['rows_by_group'].get(group_key, []))
                    if row_count == 0:
                        continue

                    cmd = build_powershell_command(script_path, temp_group_path, sheet_name)
                    process = subprocess.Popen(
                        cmd,
                        stdout=subprocess.PIPE,
                        stderr=subprocess.STDOUT,
                        universal_newlines=False
                    )
                    raw_output = process.communicate()[0]
                    process.wait()
                    full_output = decode_powershell_output(raw_output)
                    all_logs.append(f"\n[GRUPO {group_key} | {sheet_name}]\n{full_output}")

                    start_marker = '---VALIDATION_JSON_START---'
                    end_marker = '---VALIDATION_JSON_END---'
                    if start_marker not in full_output or end_marker not in full_output:
                        all_errors.append(
                            f"[GRUPO {group_key} | {sheet_name}] Não foi possível obter resultado estruturado da validação."
                        )
                        continue

                    start_idx = full_output.index(start_marker) + len(start_marker)
                    end_idx = full_output.index(end_marker)
                    json_str = full_output[start_idx:end_idx].strip()

                    try:
                        result = json.loads(json_str)
                    except Exception as exc:
                        all_errors.append(
                            f"[GRUPO {group_key} | {sheet_name}] JSON inválido da validação: {exc}"
                        )
                        continue

                    result_errors = result.get('errors') or []
                    if result.get('status') != 'success' or result_errors:
                        for err in result_errors:
                            all_errors.append(f"[GRUPO {group_key} | {sheet_name}] {err}")
            finally:
                try:
                    if os.path.exists(temp_group_path):
                        os.remove(temp_group_path)
                except Exception:
                    pass

        total_lines = total_people + total_equip
        response_payload = {
            'status': 'success' if not all_errors else 'failed',
            'total_lines': total_lines,
            'error_count': len(all_errors),
            'errors': all_errors,
            'filename': filename,
            'group_summary': group_summary,
            'sheet_counts': {
                'PESSOAS': total_people,
                'EQUIPAMENTOS': total_equip
            },
            'log': '\n'.join(all_logs)
        }
        return jsonify(response_payload), 200

    except Exception as e:
        return jsonify({
            'status': 'error',
            'errors': [f'Erro ao executar validação: {str(e)}']
        }), 500
    finally:
        end_job()


@api_bp.route('/run-script', methods=['POST'])
def run_script():
    """Fase 2: Upload para SharePoint por GRUPO (processamento em blocos)."""
    data = request.get_json()
    if not data:
        return Response('Erro: Dados não enviados.', status=400)

    filename = data.get('filename', '')
    if not filename or not allowed_file(filename):
        return Response('Erro: Arquivo inválido.', status=400)

    file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], secure_filename(filename))
    if not os.path.exists(file_path):
        return Response('Erro: Arquivo não encontrado. Execute a validação primeiro.', status=400)

    script_path = os.path.join(str(SCRIPTS_DIR), 'Populate-SharePointList.ps1')
    if not os.path.exists(script_path):
        return Response('Erro: Script Populate-SharePointList.ps1 não encontrado.', status=500)

    grouped_data, grouping_errors = _read_grouped_excel(file_path)
    if grouping_errors:
        return Response('\n'.join(grouping_errors), status=400, content_type='text/plain; charset=utf-8')

    def generate():
        start_job()
        consolidated_results: list[dict] = []

        try:
            groups = grouped_data['group_order']
            yield f'Arquivo recebido: {filename}\n'
            yield f'Processando {len(groups)} grupo(s) (PESSOAS + EQUIPAMENTOS)\n'
            yield f"{'-'*30}\n"

            stop_processing = False

            for idx, group_key in enumerate(groups, start=1):
                if stop_processing:
                    break

                people_count = len(grouped_data['sheets']['PESSOAS']['rows_by_group'].get(group_key, []))
                equip_count = len(grouped_data['sheets']['EQUIPAMENTOS']['rows_by_group'].get(group_key, []))

                yield f"\n---GROUP_PROGRESS:{idx}/{len(groups)}:{group_key}---\n"
                yield f'[GRUPO {group_key}] Iniciando processamento ({people_count} PESSOAS, {equip_count} EQUIPAMENTOS).\n'

                temp_group_path = _create_group_workbook(grouped_data, group_key)
                report_filename = None
                group_id_mob = ''
                group_ok = False

                try:
                    cmd = build_powershell_populate_command(script_path, temp_group_path)
                    process = subprocess.Popen(
                        cmd,
                        stdout=subprocess.PIPE,
                        stderr=subprocess.STDOUT,
                        bufsize=1,
                        universal_newlines=False
                    )

                    detected_error_in_output = False
                    error_markers = [
                        'FALHA CRÍTICA',
                        'UPLOAD CANCELADO',
                        '--- RESULT: ERROR ---',
                        'Write-Error',
                        'Erro ao adicionar item',
                        'Não foi possível gerar um ID_Mobilizacao único',
                    ]

                    in_report_json = False
                    report_json_lines: list[str] = []
                    report_data = None

                    REPORT_START = '---REPORT_JSON_START---'
                    REPORT_END = '---REPORT_JSON_END---'

                    while True:
                        chunk = process.stdout.readline()
                        if not chunk:
                            break
                        decoded = decode_powershell_output(chunk)
                        stripped = decoded.strip()

                        if stripped == REPORT_START:
                            in_report_json = True
                            continue
                        if stripped == REPORT_END:
                            in_report_json = False
                            continue
                        if in_report_json:
                            report_json_lines.append(stripped)
                            continue

                        if any(marker in decoded for marker in error_markers):
                            detected_error_in_output = True

                        yield decoded

                    process.wait()

                    if report_json_lines:
                        try:
                            report_data = json.loads(''.join(report_json_lines))
                            group_id_mob = str(report_data.get('id_mobilizacao', '') or '')
                            report_filename, _ = generate_report_excel(report_data)
                            yield f'---REPORT_FILE:{report_filename}---\n'
                        except Exception as exc:
                            yield f'[AVISO] Não foi possível gerar o relatório Excel do grupo {group_key}: {exc}\n'

                    group_ok = process.returncode == 0 and not detected_error_in_output
                    status_text = 'success' if group_ok else 'error'
                    group_result = {
                        'id_mobilizacao': group_id_mob,
                        'status': status_text,
                        'qtd_pessoas': people_count,
                        'qtd_equipamentos': equip_count,
                        'report_filename': report_filename
                    }
                    consolidated_results.append(group_result)
                    yield f"---GROUP_RESULT:{json.dumps(group_result, ensure_ascii=False)}---\n"

                    if group_ok:
                        yield f'[GRUPO {group_key}] Concluído com sucesso.\n'
                    else:
                        yield f'[GRUPO {group_key}] Falhou. Encerrando processamento dos próximos grupos.\n'
                        stop_processing = True

                except Exception as exc:
                    group_result = {
                        'id_mobilizacao': group_id_mob,
                        'status': 'error',
                        'qtd_pessoas': people_count,
                        'qtd_equipamentos': equip_count,
                        'report_filename': report_filename
                    }
                    consolidated_results.append(group_result)
                    yield f"---GROUP_RESULT:{json.dumps(group_result, ensure_ascii=False)}---\n"
                    yield f'[ERRO DE EXECUÇÃO][GRUPO {group_key}]: {str(exc)}\n'
                    stop_processing = True
                finally:
                    try:
                        if os.path.exists(temp_group_path):
                            os.remove(temp_group_path)
                    except Exception:
                        pass

            yield '---GROUP_SUMMARY_JSON_START---\n'
            yield json.dumps({'groups': consolidated_results}, ensure_ascii=False)
            yield '\n---GROUP_SUMMARY_JSON_END---\n'

            all_ok = consolidated_results and all(g.get('status') == 'success' for g in consolidated_results)
            if all_ok:
                yield f"\n{'-'*30}\n[SUCESSO] Todos os grupos foram processados com sucesso.\n"
            else:
                yield f"\n{'-'*30}\n[ERRO] Processamento encerrado com falha em um grupo.\n"

        except Exception as e:
            yield f'\n[ERRO DE EXECUÇÃO]: {str(e)}\n'
        finally:
            end_job()

    return Response(stream_with_context(generate()), content_type='text/plain; charset=utf-8')


@api_bp.route('/shutdown', methods=['POST'])
def shutdown():
    """Encerra o servidor Flask."""
    def _shutdown():
        os._exit(0)
    t = threading.Timer(0.5, _shutdown)
    t.daemon = True
    t.start()
    return jsonify({'message': 'Servidor encerrado.'}), 200


@api_bp.route('/list-reports')
def list_reports():
    """Lista os relatórios Excel gerados, ordenados do mais recente para o mais antigo."""
    reports_abs = os.path.abspath(current_app.config['REPORTS_FOLDER'])
    if not os.path.isdir(reports_abs):
        return jsonify({'reports': []})
    files = []
    for fname in os.listdir(reports_abs):
        if fname.lower().endswith('.xlsx'):
            fpath = os.path.join(reports_abs, fname)
            try:
                mtime = os.path.getmtime(fpath)
                size = os.path.getsize(fpath)
            except OSError:
                continue
            files.append({'name': fname, 'mtime': mtime, 'size': size})
    files.sort(key=lambda f: f['mtime'], reverse=True)
    for f in files:
        from datetime import datetime as _dt
        f['modified'] = _dt.fromtimestamp(f['mtime']).strftime('%d/%m/%Y %H:%M:%S')
        del f['mtime']
    return jsonify({'reports': files})


@api_bp.route('/download-report/<path:filename>')
def download_report(filename):
    """Serve o relatório Excel gerado após a submissão."""
    safe_name = secure_filename(filename)
    reports_abs = os.path.abspath(current_app.config['REPORTS_FOLDER'])
    file_abs = os.path.abspath(os.path.join(reports_abs, safe_name))
    if not file_abs.startswith(reports_abs + os.sep):
        return jsonify({'error': 'Acesso negado.'}), 403
    if not os.path.exists(file_abs):
        return jsonify({'error': 'Relatório não encontrado.'}), 404
    return send_from_directory(reports_abs, safe_name, as_attachment=True)


@api_bp.route('/suggest-corrections', methods=['POST'])
def suggest_corrections():
    """Recebe erros de Choice e retorna sugestões da aba LISTA SUSPENSA."""
    data = request.get_json()
    if not data:
        return jsonify({'error': 'No data'}), 400
    filename = data.get('filename', '')
    error_entries = data.get('errors', [])
    if not filename:
        return jsonify({'error': 'No filename'}), 400
    file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], secure_filename(filename))
    if not os.path.exists(file_path):
        return jsonify({'error': 'File not found'}), 404
    columns = extract_lista_suspensa_columns(file_path)
    suggestions = []
    for entry in error_entries:
        campo = entry.get('campo', '')
        valor = entry.get('valor', '')
        campo_norm = normalize_key(campo)
        matched_col = None
        for col in columns:
            if (normalize_key(col['column_name']) == campo_norm or
                    normalize_key(col['internal_name']) == campo_norm):
                matched_col = col
                break
        if not matched_col:
            for col in columns:
                if (campo_norm in normalize_key(col['column_name']) or
                        normalize_key(col['column_name']) in campo_norm or
                        campo_norm in normalize_key(col['internal_name']) or
                        normalize_key(col['internal_name']) in campo_norm):
                    matched_col = col
                    break
        if not matched_col:
            suggestions.append({**entry, 'options': [], 'suggested': None, 'column_name': campo})
            continue
        options = matched_col['options']
        valor_norm = normalize_key(valor)
        suggested = None
        best_score = 999
        for opt in options:
            opt_norm = normalize_key(opt)
            if opt_norm == valor_norm:
                suggested = opt
                break
            if valor_norm in opt_norm or opt_norm in valor_norm:
                score = abs(len(opt_norm) - len(valor_norm))
                if score < best_score:
                    best_score = score
                    suggested = opt
        suggestions.append({
            **entry,
            'options': options,
            'suggested': suggested,
            'column_name': matched_col['column_name'],
            'internal_name': matched_col['internal_name']
        })
    return jsonify({'suggestions': suggestions}), 200


@api_bp.route('/apply-corrections', methods=['POST'])
def apply_corrections():
    """Aplica correções selecionadas pelo usuário no arquivo Excel."""
    data = request.get_json()
    if not data:
        return jsonify({'error': 'No data'}), 400
    filename = data.get('filename', '')
    corrections = data.get('corrections', [])
    if not filename:
        return jsonify({'error': 'No filename'}), 400
    file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], secure_filename(filename))
    if not os.path.exists(file_path):
        return jsonify({'error': 'File not found'}), 404
    try:
        wb = openpyxl.load_workbook(file_path)
        applied = 0
        for corr in corrections:
            sheet_name = corr.get('sheet', 'PESSOAS')
            excel_row = corr.get('excel_row')
            internal_name = corr.get('internal_name', '')
            new_value = corr.get('new_value', '')
            if not excel_row or not internal_name:
                continue
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            target_col = None
            for col in range(1, ws.max_column + 1):
                header_val = ws.cell(row=1, column=col).value
                if not header_val:
                    continue
                header = str(header_val).strip()
                m = re.match(r'^(?P<disp>.+?)\s*\[(?P<internal>[^\[\]]+)\]\s*$', header)
                if m:
                    if (normalize_key(m.group('internal')) == normalize_key(internal_name) or
                            normalize_key(m.group('disp')) == normalize_key(internal_name)):
                        target_col = col
                        break
                else:
                    if normalize_key(header) == normalize_key(internal_name):
                        target_col = col
                        break
            if target_col is None:
                continue
            ws.cell(row=int(excel_row), column=target_col).value = new_value
            applied += 1
        wb.save(file_path)
        wb.close()
        return jsonify({'status': 'ok', 'applied': applied}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500
