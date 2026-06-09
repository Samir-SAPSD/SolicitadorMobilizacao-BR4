"""Report generation with Vestas styling."""

import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.config import REPORTS_FOLDER


# ── Vestas Color Palette ──
COLOR_NIGHT_SKY = '1F3144'
COLOR_BLUE_SKY = '005AFF'
COLOR_LIGHT_GREY = 'E3E5E8'
COLOR_GREEN_FILL = 'D4EDDA'
COLOR_RED_FILL = 'F8D7DA'
COLOR_WHITE = 'FFFFFF'
COLOR_LABEL_BG = 'EBF3FB'


def generate_report_excel(report_data: dict) -> tuple[str, str]:
    """
    Gera um arquivo Excel formatado a partir do JSON de relatório emitido pelo PowerShell.
    Retorna (nome_arquivo, caminho_absoluto).
    """
    fill_header = PatternFill('solid', fgColor=COLOR_NIGHT_SKY)
    fill_col_hdr = PatternFill('solid', fgColor=COLOR_NIGHT_SKY)
    fill_label = PatternFill('solid', fgColor=COLOR_LABEL_BG)
    fill_success = PatternFill('solid', fgColor=COLOR_GREEN_FILL)
    fill_error = PatternFill('solid', fgColor=COLOR_RED_FILL)
    fill_light = PatternFill('solid', fgColor=COLOR_LIGHT_GREY)

    font_white_bold = Font(name='Calibri', bold=True, color=COLOR_WHITE, size=13)
    font_col_hdr = Font(name='Calibri', bold=True, color=COLOR_WHITE, size=10)
    font_label = Font(name='Calibri', bold=True, color=COLOR_NIGHT_SKY, size=10)
    font_value = Font(name='Calibri', bold=False, color=COLOR_NIGHT_SKY, size=10)
    font_value_bold = Font(name='Calibri', bold=True, color=COLOR_NIGHT_SKY, size=10)

    thin_border_side = Side(style='thin', color='CCCCCC')
    thin_border = Border(
        left=thin_border_side, right=thin_border_side,
        top=thin_border_side, bottom=thin_border_side
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Relatório de Submissão'

    items = report_data.get('items', [])
    submission_dt = report_data.get('submission_datetime', '')
    id_mob = report_data.get('id_mobilizacao', '')
    requester_name = report_data.get('requester_name', '')
    requester_email = report_data.get('requester_email', '')

    display_cols: list[str] = []
    seen: set[str] = set()
    for item in items:
        for col_name in (item.get('fields') or {}).keys():
            if col_name not in seen:
                seen.add(col_name)
                display_cols.append(col_name)

    total_cols = 2 + len(display_cols)
    last_col_letter = get_column_letter(max(total_cols, 3))

    ws.merge_cells(f'A1:{last_col_letter}1')
    title_cell = ws['A1']
    title_cell.value = 'RELATÓRIO DE SUBMISSÃO — MOBILIZAÇÕES VESTAS'
    title_cell.fill = fill_header
    title_cell.font = font_white_bold
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    meta_rows = [
        ('Data/Hora da Submissão:', submission_dt),
        ('ID de Mobilização:', id_mob),
        ('Solicitante:', requester_name),
        ('E-mail:', requester_email),
    ]
    for r_offset, (label, value) in enumerate(meta_rows, start=2):
        label_cell = ws.cell(row=r_offset, column=1, value=label)
        label_cell.fill = fill_label
        label_cell.font = font_label
        label_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        label_cell.border = thin_border

        value_cell = ws.cell(row=r_offset, column=2, value=value)
        value_cell.fill = fill_light
        value_cell.font = font_value_bold
        value_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        value_cell.border = thin_border

        if total_cols > 2:
            ws.merge_cells(start_row=r_offset, start_column=2, end_row=r_offset, end_column=total_cols)
        ws.row_dimensions[r_offset].height = 18

    ws.row_dimensions[7].height = 8

    hdr_row = 8
    headers = ['ID Elemento', 'Status'] + display_cols
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=hdr_row, column=col_idx, value=header_text)
        cell.fill = fill_col_hdr
        cell.font = font_col_hdr
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[hdr_row].height = 22

    for row_offset, item in enumerate(items, start=hdr_row + 1):
        is_error = 'Erro' in str(item.get('status', ''))
        row_fill = fill_error if is_error else fill_success

        id_sp_cell = ws.cell(row=row_offset, column=1, value=item.get('id_sp', ''))
        id_sp_cell.fill = row_fill
        id_sp_cell.font = font_value_bold
        id_sp_cell.alignment = Alignment(horizontal='center', vertical='center')
        id_sp_cell.border = thin_border

        status_cell = ws.cell(row=row_offset, column=2, value=item.get('status', ''))
        status_cell.fill = row_fill
        status_cell.font = font_value
        status_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)
        status_cell.border = thin_border

        fields = item.get('fields') or {}
        for col_offset, col_name in enumerate(display_cols, start=3):
            val = fields.get(col_name, '')
            data_cell = ws.cell(row=row_offset, column=col_offset, value=val)
            data_cell.fill = row_fill
            data_cell.font = font_value
            data_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)
            data_cell.border = thin_border

        ws.row_dimensions[row_offset].height = 18

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 30
    for col_offset, col_name in enumerate(display_cols, start=3):
        estimated = max(14, min(len(col_name) + 4, 40))
        ws.column_dimensions[get_column_letter(col_offset)].width = estimated

    ws.freeze_panes = f'A{hdr_row + 1}'

    os.makedirs(REPORTS_FOLDER, exist_ok=True)
    safe_id = re.sub(r'[^A-Za-z0-9]', '_', id_mob) if id_mob else 'sem_id'
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    report_filename = f'Relatorio_Mob_{safe_id}_{timestamp}.xlsx'
    report_path = os.path.abspath(os.path.join(REPORTS_FOLDER, report_filename))
    wb.save(report_path)
    return report_filename, report_path
