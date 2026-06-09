"""Excel file reading and processing utilities."""

import os
import tempfile

import openpyxl

from backend.config import REPORTS_FOLDER


def read_grouped_excel(file_path: str) -> tuple[dict | None, list[str]]:
    """
    Lê abas PESSOAS e EQUIPAMENTOS exigindo coluna A = GRUPO.
    Retorna (data, errors), onde data contém headers, linhas por grupo e ordem de grupos.
    """
    errors: list[str] = []
    required_sheets = ['PESSOAS', 'EQUIPAMENTOS']

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as exc:
        return None, [f'Não foi possível abrir o Excel: {exc}']

    try:
        for sheet in required_sheets:
            if sheet not in wb.sheetnames:
                errors.append(f"A aba obrigatória '{sheet}' não foi encontrada.")

        if errors:
            return None, errors

        groups_order: list[str] = []
        groups_seen: set[str] = set()

        data = {
            'group_order': groups_order,
            'sheets': {
                'PESSOAS': {'headers': [], 'rows': [], 'rows_by_group': {}},
                'EQUIPAMENTOS': {'headers': [], 'rows': [], 'rows_by_group': {}},
            }
        }

        for sheet in required_sheets:
            ws = wb[sheet]
            max_col = max(1, ws.max_column)
            headers = [ws.cell(row=1, column=col).value for col in range(1, max_col + 1)]
            data['sheets'][sheet]['headers'] = headers

            header_a = '' if headers[0] is None else str(headers[0]).strip()
            if header_a != 'GRUPO':
                errors.append(
                    f"A aba '{sheet}' deve ter a coluna A com cabeçalho exatamente 'GRUPO'."
                )
                continue

            duplicate_keys: dict[str, set[tuple[str, ...]]] = {}

            for row_idx in range(2, ws.max_row + 1):
                row_values = [ws.cell(row=row_idx, column=col).value for col in range(1, max_col + 1)]
                if not row_has_any_data(row_values):
                    continue

                group_raw = row_values[0]
                if group_raw is None or str(group_raw).strip() == '':
                    errors.append(f'Aba {sheet}, linha {row_idx}: GRUPO vazio.')
                    continue

                group_key = value_to_exact_text(group_raw)
                signature = tuple(value_to_exact_text(v) for v in row_values)

                duplicate_bucket = duplicate_keys.setdefault(group_key, set())
                if signature in duplicate_bucket:
                    errors.append(
                        f"Aba {sheet}, linha {row_idx}: linha duplicada dentro do GRUPO '{group_key}'."
                    )
                    continue
                duplicate_bucket.add(signature)

                row_data = {
                    'excel_row': row_idx,
                    'group': group_key,
                    'values': row_values,
                }
                data['sheets'][sheet]['rows'].append(row_data)
                data['sheets'][sheet]['rows_by_group'].setdefault(group_key, []).append(row_data)

                if group_key not in groups_seen:
                    groups_seen.add(group_key)
                    groups_order.append(group_key)

        if errors:
            return None, errors

        if not groups_order:
            return None, [
                'Nenhuma linha válida encontrada para importar no SharePoint. As abas PESSOAS e EQUIPAMENTOS estão vazias.'
            ]

        return data, []
    finally:
        try:
            wb.close()
        except Exception:
            pass


def create_group_workbook(source_data: dict, group_key: str) -> str:
    """Cria um arquivo temporário .xlsx contendo apenas linhas do grupo informado."""
    out_wb = openpyxl.Workbook()

    default_sheet = out_wb.active
    out_wb.remove(default_sheet)

    for sheet_name in ['PESSOAS', 'EQUIPAMENTOS']:
        ws = out_wb.create_sheet(title=sheet_name)
        headers = source_data['sheets'][sheet_name]['headers']
        ws.append(headers)
        group_rows = source_data['sheets'][sheet_name]['rows_by_group'].get(group_key, [])
        for row in group_rows:
            ws.append(row['values'])

    os.makedirs(REPORTS_FOLDER, exist_ok=True)
    temp_fd, temp_path = tempfile.mkstemp(prefix='mob_group_', suffix='.xlsx', dir=REPORTS_FOLDER)
    os.close(temp_fd)
    out_wb.save(temp_path)
    out_wb.close()
    return temp_path


def count_excel_rows(file_path: str, sheet_name: str) -> int:
    """Conta linhas com dados (excluindo cabeçalho) em uma aba do Excel usando openpyxl."""
    wb = None
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            return 0
        ws = wb[sheet_name]
        count = 0
        for row in ws.iter_rows(min_row=2):
            if any(cell.value is not None and str(cell.value).strip() for cell in row):
                count += 1
        return count
    except Exception:
        return 0
    finally:
        try:
            if wb is not None:
                wb.close()
        except Exception:
            pass


def value_to_exact_text(value) -> str:
    """Converte valor de célula para comparação textual exata (sem normalização)."""
    if value is None:
        return ''
    return str(value)


def row_has_any_data(values: list) -> bool:
    """Verifica se alguma célula da linha tem dados."""
    for value in values:
        if value is None:
            continue
        if str(value).strip() != '':
            return True
    return False
