"""Helpers for Excel correction suggestions and validation."""

import re
import unicodedata

import openpyxl


def normalize_key(text: str) -> str:
    """Normalize text: strip, lower, remove accents for comparison."""
    if not text:
        return ''
    nfd = unicodedata.normalize('NFD', str(text).strip().lower())
    return ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')


def extract_lista_suspensa_columns(file_path: str) -> list:
    """
    Reads the hidden 'LISTA SUSPENSA' sheet from an Excel file.
    Returns list of dicts: [{column_name, internal_name, options: [str]}]
    """
    result = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet_name = None
        for name in wb.sheetnames:
            if normalize_key(name) in ('lista suspensa', 'lista_suspensa', 'listasuspensa', 'dropdown'):
                sheet_name = name
                break
        if not sheet_name:
            wb.close()
            return result
        ws = wb[sheet_name]
        max_col = ws.max_column
        if not max_col:
            wb.close()
            return result
        for col in range(1, max_col + 1):
            header_cell = ws.cell(row=1, column=col).value
            if not header_cell:
                continue
            header = str(header_cell).strip()
            m = re.match(r'^(?P<disp>.+?)\s*\[(?P<internal>[^\[\]]+)\]\s*$', header)
            if m:
                display_name = m.group('disp').strip()
                internal_name = m.group('internal').strip()
            else:
                display_name = header
                internal_name = header
            options = []
            for row in range(2, ws.max_row + 1):
                val = ws.cell(row=row, column=col).value
                if val is not None and str(val).strip():
                    options.append(str(val).strip())
            if options:
                result.append({
                    'column_name': display_name,
                    'internal_name': internal_name,
                    'options': options
                })
        wb.close()
    except Exception:
        pass
    return result


def allowed_file(filename: str, allowed_extensions: set = None) -> bool:
    """Verifica se filename tem uma extensão permitida."""
    if allowed_extensions is None:
        allowed_extensions = {'xlsx'}
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in allowed_extensions
